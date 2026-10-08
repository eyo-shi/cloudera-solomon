"""IDアドレスマップ (Excel) を用いた hex ダンプのフィールド解釈。"""
from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal

from solomon.graph.manufacturing.models import DumpFieldNode
from solomon.graph.manufacturing.id_dump import parse_hex_dump_lines, read_bytes_at

TimingFilter = Literal["main_out", "shukken_out"]

_MAIN_TIMING = frozenset({"ﾒｲﾝｱｳﾄ", "メインアウト", "MAIN"})
_SHUKKEN_TIMING = frozenset({"出検ｱｳﾄ", "出検アウト", "出検ｱｳﾄシュッケン", "SHUKKEN"})

_JUDGE_BCD = {
    "00": "未検",
    "01": "OK",
    "02": "NG",
    "03": "手直しOK",
}

# T447 ID メモリ上の代表 composite (ヘッダ KIBAN/HINBAN と突合)
_COMPOSITE_FIELDS: tuple[tuple[str, int, int], ...] = (
    ("KIBAN", 0x0036, 10),
    ("HINBAN", 0x0016, 10),
)

_ADDR_RE = re.compile(r"^([0-9A-Fa-f]{4})\(H\)$")
_CONTINUATION_MARKERS = frozenset({"↑", "↑ ↑"})


@dataclass(frozen=True)
class AddressMapEntry:
    address: int
    address_hex: str
    info: str
    timing: str
    scale: str = ""
    process_name: str = ""


@dataclass(frozen=True)
class FieldSpan:
    start: int
    end: int
    info: str
    timing: str
    scale: str
    process_name: str


def resolve_address_map_path(explicit: str | None = None) -> Path | None:
    """IDアドレスマップ Excel/JSON のパスを解決する。"""
    candidates: list[str] = []
    if explicit:
        candidates.append(explicit)
    env_path = os.environ.get("SOLOMON_ID_ADDRESS_MAP_XLSX") or os.environ.get(
        "SOLOMON_ID_ADDRESS_MAP"
    )
    if env_path:
        candidates.append(env_path)
    candidates.extend(
        [
            "demo_data/received/id_address_map_t447.json",
            "demo_data/received/id_address_map_t447.xlsx",
        ]
    )
    for candidate in candidates:
        path = Path(candidate).expanduser()
        if path.is_file():
            return path
    default_xlsx = Path(
        "/Users/eyoshida/dev/file_sample/"
        "【IDマスタ】QMSD22301_T447ライン追加設定設定確認書_IDアドレスマップ.xlsx"
    )
    if default_xlsx.is_file():
        return default_xlsx
    fixture = Path(__file__).resolve().parents[3] / "tests/fixtures/id_address_map_t447_subset.json"
    if fixture.is_file():
        return fixture
    return None


def timing_for_data_type(data_type: str) -> TimingFilter | None:
    if data_type == "j5_main_id":
        return "main_out"
    if data_type == "j5_shukken_id":
        return "shukken_out"
    return None


def _normalize_timing(timing: str) -> str:
    return timing.strip()


def _matches_timing(timing: str, target: TimingFilter) -> bool:
    normalized = _normalize_timing(timing)
    if target == "main_out":
        return normalized in _MAIN_TIMING
    return normalized in _SHUKKEN_TIMING


def load_address_map_entries(path: Path) -> list[AddressMapEntry]:
    """Excel または JSON からアドレスマップ行を読み込む。"""
    if path.suffix.lower() == ".json":
        return _load_entries_from_json(path)
    return _load_entries_from_xlsx(path)


def _load_entries_from_json(path: Path) -> list[AddressMapEntry]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    entries: list[AddressMapEntry] = []
    for row in raw:
        addr_hex = str(row.get("address_hex", "")).upper()
        if not _ADDR_RE.match(addr_hex):
            continue
        entries.append(
            AddressMapEntry(
                address=int(addr_hex[:4], 16),
                address_hex=addr_hex,
                info=str(row.get("info", "")).strip(),
                timing=str(row.get("timing", "")).strip(),
                scale=str(row.get("scale", "")),
                process_name=str(row.get("process_name", "")),
            )
        )
    return sorted(entries, key=lambda item: item.address)


def _load_entries_from_xlsx(path: Path) -> list[AddressMapEntry]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:  # pragma: no cover - optional at runtime
        raise RuntimeError("openpyxl is required to load ID address map xlsx") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = "IDアドレスマップ設定確認書"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"sheet not found: {sheet_name}")
    sheet = workbook[sheet_name]
    entries: list[AddressMapEntry] = []
    for row_index, row in enumerate(sheet.iter_rows(min_row=8, values_only=True)):
        if not row or len(row) < 23:
            continue
        addr = row[5]
        if not isinstance(addr, str) or not _ADDR_RE.match(addr.strip()):
            continue
        addr_hex = addr.strip().upper()
        info = str(row[6] or "").strip()
        timing = str(row[22] or "").strip()
        entries.append(
            AddressMapEntry(
                address=int(addr_hex[:4], 16),
                address_hex=addr_hex,
                info=info,
                timing=timing,
                scale=str(row[25] or ""),
                process_name=str(row[26] or ""),
            )
        )
    workbook.close()
    return sorted(entries, key=lambda item: item.address)


def group_field_spans(entries: list[AddressMapEntry]) -> list[FieldSpan]:
    """同一フィールド (情報内容 + ↑ 連続) をアドレススパンにまとめる。"""
    spans: list[FieldSpan] = []
    current: FieldSpan | None = None
    for entry in entries:
        info = entry.info.strip()
        if info in _CONTINUATION_MARKERS or info.startswith("↑"):
            if current is not None:
                current = FieldSpan(
                    start=current.start,
                    end=entry.address,
                    info=current.info,
                    timing=current.timing,
                    scale=current.scale,
                    process_name=current.process_name,
                )
            continue
        if not info:
            continue
        if current is not None:
            spans.append(current)
        current = FieldSpan(
            start=entry.address,
            end=entry.address,
            info=info,
            timing=entry.timing,
            scale=entry.scale,
            process_name=entry.process_name,
        )
    if current is not None:
        spans.append(current)
    return spans


def decode_span_value(raw_bytes: list[int], scale: str) -> tuple[str, str]:
    """スパンの生バイト列を表示値に変換する。"""
    if not raw_bytes:
        return "", "empty"

    scale_upper = (scale or "").upper()
    if "BCD" in scale_upper and "F:" in scale_upper:
        code = f"{raw_bytes[0]:02X}"
        if len(raw_bytes) == 1:
            bcd = f"{raw_bytes[0]:02d}"
            return _JUDGE_BCD.get(bcd, bcd), "judge_bcd"
        return code, "bcd"

    ascii_chars: list[str] = []
    for byte in raw_bytes:
        if byte in (0x00, 0x20):
            continue
        if 32 <= byte < 127:
            ascii_chars.append(chr(byte))
    if ascii_chars:
        return "".join(ascii_chars), "ascii"

    if len(raw_bytes) == 1:
        return f"{raw_bytes[0]:02X}", "hex"
    return ",".join(f"{byte:02X}" for byte in raw_bytes), "hex"


@lru_cache(maxsize=2)
def _cached_map_entries(path_str: str, mtime_ns: int) -> tuple[AddressMapEntry, ...]:
    path = Path(path_str)
    entries = load_address_map_entries(path)
    return tuple(entries)


def get_address_map_entries(path: Path | None = None) -> list[AddressMapEntry]:
    resolved = path or resolve_address_map_path()
    if resolved is None:
        return []
    stat = resolved.stat()
    return list(_cached_map_entries(str(resolved.resolve()), stat.st_mtime_ns))


def decode_composite_fields(
    memory: dict[int, int],
    *,
    trace_id: str = "",
    timing: str = "",
) -> list[DumpFieldNode]:
    """アドレスマップ外の代表 composite (KIBAN/HINBAN 等) を解釈する。"""
    fields: list[DumpFieldNode] = []
    for name, start, length in _COMPOSITE_FIELDS:
        raw_bytes = read_bytes_at(memory, start, length)
        value, value_type = decode_span_value(raw_bytes, "ascii")
        if not value:
            continue
        end = start + length - 1
        fields.append(
            DumpFieldNode(
                address_hex=f"{start:04X}(H)",
                address_end_hex=f"{end:04X}(H)",
                label=name,
                raw_hex=",".join(f"{byte:02X}" for byte in raw_bytes),
                value=value,
                value_type=value_type,
                timing=timing,
                process_name="composite",
                trace_id=trace_id,
            )
        )
    return fields


def decode_id_dump_fields(
    text: str,
    *,
    data_type: str,
    map_path: str | Path | None = None,
    trace_id: str = "",
) -> list[DumpFieldNode]:
    """hex ダンプテキストを ID アドレスマップで解釈する。"""
    timing = timing_for_data_type(data_type)
    if timing is None:
        return []

    path = Path(map_path) if map_path else resolve_address_map_path()
    if path is None:
        return []

    entries = [
        entry
        for entry in get_address_map_entries(path)
        if _matches_timing(entry.timing, timing)
    ]
    if not entries:
        return []

    memory = parse_hex_dump_lines(text)
    if not memory:
        return []

    timing_label = entries[0].timing if entries else ""
    fields: list[DumpFieldNode] = decode_composite_fields(
        memory, trace_id=trace_id, timing=timing_label
    )
    seen_ids = {field.id for field in fields}

    for span in group_field_spans(entries):
        length = span.end - span.start + 1
        raw_bytes = read_bytes_at(memory, span.start, length)
        if not any(raw_bytes):
            continue
        value, value_type = decode_span_value(raw_bytes, span.scale)
        if not value:
            continue
        start_hex = f"{span.start:04X}(H)"
        end_hex = f"{span.end:04X}(H)" if span.end != span.start else ""
        node = DumpFieldNode(
            address_hex=start_hex,
            address_end_hex=end_hex,
            label=span.info,
            raw_hex=",".join(f"{byte:02X}" for byte in raw_bytes),
            value=value,
            value_type=value_type,
            timing=span.timing,
            process_name=span.process_name,
            trace_id=trace_id,
        )
        if node.id in seen_ids:
            continue
        seen_ids.add(node.id)
        fields.append(node)
    return fields
